import re
import sys
import os

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CBD5E1')
    pBdr.append(bottom)
    pPr.append(pBdr)

def render_inline(paragraph, text, base_font_size=10.5, base_color=RGBColor(0x33, 0x41, 0x55), base_bold=False):
    tokens = []
    # Tokenize bold (**...**), italic (*...*), and inline code (`...`)
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    last_idx = 0
    for match in pattern.finditer(text):
        if match.start() > last_idx:
            tokens.append(('text', text[last_idx:match.start()]))
        raw = match.group(0)
        if raw.startswith('**') and raw.endswith('**'):
            tokens.append(('bold', raw[2:-2]))
        elif raw.startswith('*') and raw.endswith('*'):
            tokens.append(('italic', raw[1:-1]))
        elif raw.startswith('`') and raw.endswith('`'):
            tokens.append(('code', raw[1:-1]))
        last_idx = match.end()
    if last_idx < len(text):
        tokens.append(('text', text[last_idx:]))

    for token_type, val in tokens:
        run = paragraph.add_run(val)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(base_font_size)
        run.font.color.rgb = base_color
        run.bold = base_bold

        if token_type == 'bold':
            run.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif token_type == 'italic':
            run.italic = True
        elif token_type == 'code':
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size - 1)
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

def convert_markdown_to_docx(input_path, output_path):
    print(f"Reading Markdown from: {input_path}")
    with open(input_path, 'r', encoding='utf-8', errors='replace') as f:
        md_text = f.read()

    doc = docx.Document()

    # Set 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Configure Normal style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Segoe UI'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    lines = md_text.splitlines()
    in_code_block = False
    code_block_lines = []

    for line in lines:
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            if in_code_block:
                code_text = '\n'.join(code_block_lines)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                code_block_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_block_lines = []
            continue

        if in_code_block:
            code_block_lines.append(line)
            continue

        # Blank line
        if not stripped:
            continue

        # Horizontal rule
        if stripped in ('---', '***', '___'):
            add_horizontal_rule(doc)
            continue

        # Heading 1 (# ...)
        if stripped.startswith('# '):
            text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Segoe UI'
            run.font.size = Pt(20)
            run.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Dark Navy/Slate
            continue

        # Heading 2 (## ...)
        if stripped.startswith('## '):
            text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Segoe UI'
            run.font.size = Pt(14)
            run.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            continue

        # Heading 3 (### ...)
        if stripped.startswith('### '):
            text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Segoe UI'
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            continue

        # Checkbox list (- [ ] or - [x])
        chk_match = re.match(r'^[-*]\s*\[([ xX])\]\s*(.*)$', stripped)
        if chk_match:
            is_checked = chk_match.group(1).lower() == 'x'
            item_text = chk_match.group(2)
            symbol = '☑  ' if is_checked else '☐  '
            symbol_color = RGBColor(0x10, 0xB9, 0x81) if is_checked else RGBColor(0x64, 0x74, 0x8B)

            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)

            sym_run = p.add_run(symbol)
            sym_run.font.name = 'Segoe UI Symbol'
            sym_run.font.size = Pt(12)
            sym_run.bold = True
            sym_run.font.color.rgb = symbol_color

            render_inline(p, item_text, base_font_size=10.5)
            continue

        # Bullet list (- or *)
        if stripped.startswith(('- ', '* ')):
            item_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)

            bullet_run = p.add_run('•  ')
            bullet_run.font.name = 'Segoe UI'
            bullet_run.font.size = Pt(11)
            bullet_run.bold = True
            bullet_run.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5) # Indigo bullet

            render_inline(p, item_text, base_font_size=10.5)
            continue

        # Timestamped speaker lines, e.g.: [00:00] Speaker 1: "..."
        ts_match = re.match(r'^(\[\d{1,2}:\d{2}(?::\d{2})?\])\s*([^:]+):\s*(.*)$', stripped)
        if ts_match:
            ts = ts_match.group(1)
            speaker = ts_match.group(2)
            speech = ts_match.group(3)

            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)

            # Timestamp
            ts_run = p.add_run(f"{ts} ")
            ts_run.font.name = 'Segoe UI'
            ts_run.font.size = Pt(9.5)
            ts_run.bold = True
            ts_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

            # Speaker
            spk_run = p.add_run(f"{speaker}: ")
            spk_run.font.name = 'Segoe UI'
            spk_run.font.size = Pt(10.5)
            spk_run.bold = True
            spk_run.font.color.rgb = RGBColor(0x43, 0x38, 0xCA) # Deep Indigo

            # Speech
            render_inline(p, speech, base_font_size=10.5)
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        render_inline(p, stripped, base_font_size=10.5)

    doc.save(output_path)
    print(f"✅ Successfully converted to DOCX: {output_path}")

if __name__ == '__main__':
    in_file = sys.argv[1] if len(sys.argv) > 1 else r"c:\LOCAL_DISK_(D)\Projects\infotech_Meetings\Meeting_Summary_2026-09-22_bjp-gsuf-cvh.md"
    out_file = sys.argv[2] if len(sys.argv) > 2 else in_file.rsplit('.', 1)[0] + ".docx"
    convert_markdown_to_docx(in_file, out_file)
