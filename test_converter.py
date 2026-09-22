import os
import docx
import convert_to_docx

def test_full_pipeline():
    test_md = """# Test Executive Meeting
## Key Decisions
- [ ] Task 1 with **bold text** and `inline code`
- [x] Task 2 completed
- Standard bullet with *italics*

---

| Feature | Status | Notes |
| --- | --- | --- |
| Docx Sync | Done | Native Word XML |
| Resumable Upload | Done | Bypasses Apps Script limits |

[00:00] Speaker 1: "The meeting is officially underway."
[01:30] Speaker 2: "Confirmed. All systems are operational."
"""
    md_file = "temp_test_meeting.md"
    docx_file = "temp_test_meeting.docx"

    with open(md_file, "w", encoding="utf-8") as f:
        f.write(test_md)

    try:
        convert_to_docx.convert_markdown_to_docx(md_file, docx_file)
        assert os.path.exists(docx_file), "Output docx must exist"
        size = os.path.getsize(docx_file)
        assert size > 2000, f"Docx must be non-empty (got {size} bytes)"

        doc = docx.Document(docx_file)
        assert len(doc.paragraphs) > 5, "Must contain styled paragraphs"
        assert len(doc.tables) == 1, "Must contain parsed table"
        assert len(doc.tables[0].rows) == 3, "Table must have 3 rows"
        print(f"✅ PASS: Python DOCX Converter verified successfully ({len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {size} bytes)")
    finally:
        if os.path.exists(md_file): os.remove(md_file)
        if os.path.exists(docx_file): os.remove(docx_file)

if __name__ == "__main__":
    test_full_pipeline()
